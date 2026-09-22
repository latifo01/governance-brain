from __future__ import annotations

import shutil
import subprocess
import tempfile
import zipfile
import os
from collections.abc import Iterable
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from gov360_brain.contracts import (
    NormalizationContext,
    NormalizationResult,
    ProbeResult,
    SourceInventory,
)
from gov360_brain.utils import atomic_write_bytes, escape_markdown_cell, safe_filename

from .base import (
    DependencyUnavailableError,
    QuarantinedSourceError,
    ValidatingAdapter,
    markdown_table,
    neutralize_untrusted_text,
    unit,
)
from .registry import validate_safe_zip


OOXML_SIGNATURE = b"PK\x03\x04"
OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")


def _powerpoint_renders(path: Path, output: Path) -> tuple[list[dict[str, Any]], str | None]:
    """Render from a temporary copy with macros disabled; never touch the source."""

    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return [], "PowerPoint rendering is unavailable; install the local Windows Office extra"
    stage = Path(tempfile.mkdtemp(prefix="gov360-ppt-render-"))
    application = presentation = None
    pythoncom.CoInitialize()
    try:
        copied = stage / path.name
        shutil.copy2(path, copied)
        application = win32com.client.DispatchEx("PowerPoint.Application")
        application.AutomationSecurity = 3
        presentation = application.Presentations.Open(str(copied), True, False, False)
        rendered = stage / "slides"
        presentation.Export(str(rendered), "PNG")
        artifacts: list[dict[str, Any]] = []
        for index, item in enumerate(sorted(rendered.glob("*.PNG")), 1):
            target = output / "data" / f"slide-{index:04d}.png"
            atomic_write_bytes(target, item.read_bytes())
            artifacts.append({"artifact_type": "slide_render", "relative_path": f"data/{target.name}", "locator": f"Slide {index}"})
        return artifacts, None
    except Exception:
        return [], "PowerPoint local rendering failed; visual review remains required"
    finally:
        if presentation is not None:
            try:
                presentation.Close()
            except Exception:
                pass
        if application is not None:
            try:
                application.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()
        shutil.rmtree(stage, ignore_errors=True)


def _excel_chart_renders(path: Path, output: Path) -> tuple[list[dict[str, Any]], str | None]:
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return [], "Excel chart rendering is unavailable; install the local Windows Office extra"
    stage = Path(tempfile.mkdtemp(prefix="gov360-xls-render-"))
    application = workbook = None
    pythoncom.CoInitialize()
    try:
        copied = stage / path.name
        shutil.copy2(path, copied)
        application = win32com.client.DispatchEx("Excel.Application")
        application.Visible = False
        application.DisplayAlerts = False
        application.EnableEvents = False
        application.AskToUpdateLinks = False
        application.AutomationSecurity = 3
        application.Calculation = -4135
        workbook = application.Workbooks.Open(str(copied), UpdateLinks=0, ReadOnly=True, IgnoreReadOnlyRecommended=True, AddToMru=False)
        artifacts: list[dict[str, Any]] = []
        number = 0
        for sheet in workbook.Worksheets:
            chart_objects = sheet.ChartObjects()
            for index in range(1, chart_objects.Count + 1):
                number += 1
                chart = chart_objects.Item(index)
                staged = stage / f"chart-{number:04d}.png"
                chart.Chart.Export(str(staged), "PNG")
                target = output / "data" / staged.name
                atomic_write_bytes(target, staged.read_bytes())
                artifacts.append({"artifact_type": "chart_render", "relative_path": f"data/{target.name}", "locator": f"Sheet {sheet.Name}, chart {index}"})
        return artifacts, None
    except Exception:
        return [], "Excel local chart rendering failed; visual review remains required"
    finally:
        if workbook is not None:
            try:
                workbook.Close(False)
            except Exception:
                pass
        if application is not None:
            try:
                application.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()
        shutil.rmtree(stage, ignore_errors=True)


def _zip_names(path: Path) -> set[str]:
    validate_safe_zip(path)
    with zipfile.ZipFile(path) as package:
        return set(package.namelist())


def _package_xml(path: Path, member: str) -> bytes:
    validate_safe_zip(path)
    with zipfile.ZipFile(path) as package:
        try:
            return package.read(member)
        except KeyError as exc:
            raise QuarantinedSourceError("required document package part is missing") from exc


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _ooxml_probe(path: Path, *, root: str, extensions: set[str], fmt: str, mime: str) -> ProbeResult | None:
    with path.open("rb") as stream:
        if stream.read(4) != OOXML_SIGNATURE:
            return None
    try:
        names = _zip_names(path)
    except QuarantinedSourceError:
        raise
    if not any(name.startswith(root + "/") for name in names):
        return None
    warnings: tuple[str, ...] = ()
    if path.suffix.lower() not in extensions:
        warnings = ("extension does not match the detected Office package",)
    return ProbeResult(fmt, fmt, mime, 1.0, warnings)


class WordProcessingAdapter(ValidatingAdapter):
    adapter_id = "word-processing"
    capabilities = {"air_gap": True, "native_text": True, "rendering": "local-office", "ocr": False, "tables": True, "charts": False}

    def probe(self, path: Path) -> ProbeResult | None:
        result = _ooxml_probe(
            path,
            root="word",
            extensions={".docx", ".docm"},
            fmt="docx" if path.suffix.lower() != ".docm" else "docm",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        if result is None:
            return None
        return ProbeResult(self.adapter_id, result.source_format, result.mime_type, result.confidence, result.warnings)

    @staticmethod
    def _document(path: Path):
        try:
            from docx import Document
        except ImportError as exc:
            raise DependencyUnavailableError("python-docx is required for Word normalization") from exc
        try:
            return Document(path)
        except Exception as exc:
            raise QuarantinedSourceError("Word package cannot be read") from exc

    def inventory(self, path: Path, context: NormalizationContext | None = None) -> SourceInventory:
        document = self._document(path)
        names = _zip_names(path)
        metadata = {
            "paragraph_count": len(document.paragraphs),
            "table_count": len(document.tables),
            "section_count": len(document.sections),
            "hyperlink_relationship_count": sum(
                1 for rel in document.part.rels.values() if "hyperlink" in rel.reltype
            ),
            "comments_present": "word/comments.xml" in names,
            "footnotes_present": "word/footnotes.xml" in names,
            "endnotes_present": "word/endnotes.xml" in names,
            "macros_present": "word/vbaProject.bin" in names,
        }
        return SourceInventory("docm" if metadata["macros_present"] else "docx", "document", max(1, len(document.paragraphs) + len(document.tables)), metadata)

    @staticmethod
    def _blocks(document: Any) -> Iterable[tuple[str, Any]]:
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P

        for child in document.element.body.iterchildren():
            if isinstance(child, CT_P):
                yield "paragraph", Paragraph(child, document)
            elif isinstance(child, CT_Tbl):
                yield "table", Table(child, document)

    def normalize(self, path: Path, context: NormalizationContext) -> NormalizationResult:
        document = self._document(path)
        inventory = self.inventory(path)
        blocks: list[str] = []
        for kind, value in self._blocks(document):
            if kind == "paragraph":
                style = str(getattr(getattr(value, "style", None), "name", "") or "")
                prefix = "### " if style.lower().startswith("heading") else ""
                blocks.append(prefix + neutralize_untrusted_text(value.text))
            else:
                rows = [[cell.text for cell in row.cells] for row in value.rows]
                if rows:
                    blocks.append(markdown_table(rows[0], rows[1:]))
        if not blocks:
            blocks = ["> No extractable native text or tables were found."]

        chunks: list[tuple[int, int, str]] = []
        start = 1
        current: list[str] = []
        size = 0
        for number, block in enumerate(blocks, 1):
            if current and size + len(block) + 2 > context.max_markdown_chars:
                chunks.append((start, number - 1, "\n\n".join(current)))
                current, size, start = [], 0, number
            current.append(block)
            size += len(block) + 2
        if current:
            chunks.append((start, len(blocks), "\n\n".join(current)))

        warnings: list[str] = []
        for key in ("comments_present", "footnotes_present", "endnotes_present"):
            if inventory.metadata.get(key):
                warnings.append(f"{key.replace('_', ' ')} require review because the native reader has limited coverage")
        if inventory.metadata.get("macros_present"):
            warnings.append("macros were detected and not executed")
        units = [
            unit(
                locator_kind="section",
                locator=f"Body blocks {first}-{last}",
                filename=f"section-{index:04d}.md",
                body=f"## Body blocks {first}-{last}\n\n{body}",
                metadata={"block_start": first, "block_end": last, "extraction_mode": "native"},
                warnings=list(warnings),
            )
            for index, (first, last, body) in enumerate(chunks, 1)
        ]
        return NormalizationResult(
            self.adapter_id,
            self.adapter_version,
            inventory.source_format,
            "document",
            inventory,
            units,
            warnings=warnings,
            metrics={"native_block_count": len(blocks), "unit_count": len(units)},
        )


class PresentationAdapter(ValidatingAdapter):
    adapter_id = "presentation"
    heavy = True
    capabilities = {"air_gap": True, "native_text": True, "rendering": "local-office", "ocr": False, "tables": True, "charts": True}

    def probe(self, path: Path) -> ProbeResult | None:
        fmt = "pptm" if path.suffix.lower() == ".pptm" else "pptx"
        result = _ooxml_probe(
            path,
            root="ppt",
            extensions={".pptx", ".pptm"},
            fmt=fmt,
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
        if result is None:
            return None
        return ProbeResult(self.adapter_id, result.source_format, result.mime_type, result.confidence, result.warnings)

    @staticmethod
    def _presentation(path: Path):
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise DependencyUnavailableError("python-pptx is required for presentation normalization") from exc
        try:
            return Presentation(path)
        except Exception as exc:
            raise QuarantinedSourceError("presentation package cannot be read") from exc

    @staticmethod
    def _walk_shapes(shapes: Any) -> Iterable[Any]:
        for shape in shapes:
            yield shape
            if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):
                yield from PresentationAdapter._walk_shapes(shape.shapes)

    def inventory(self, path: Path, context: NormalizationContext | None = None) -> SourceInventory:
        presentation = self._presentation(path)
        names = _zip_names(path)
        chart_count = picture_count = table_count = note_count = hyperlink_count = 0
        for slide in presentation.slides:
            for shape in self._walk_shapes(slide.shapes):
                chart_count += int(bool(getattr(shape, "has_chart", False)))
                table_count += int(bool(getattr(shape, "has_table", False)))
                picture_count += int(getattr(shape, "shape_type", None) == 13)
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        hyperlink_count += sum(int(bool(run.hyperlink.address)) for run in paragraph.runs)
            try:
                note_count += int(bool(slide.notes_slide.notes_text_frame.text.strip()))
            except Exception:
                pass
        metadata = {
            "slide_count": len(presentation.slides),
            "chart_count": chart_count,
            "picture_count": picture_count,
            "table_count": table_count,
            "notes_slide_count": note_count,
            "hyperlink_count": hyperlink_count,
            "macros_present": "ppt/vbaProject.bin" in names,
        }
        return SourceInventory("pptm" if metadata["macros_present"] else "pptx", "presentation", len(presentation.slides), metadata)

    def normalize(self, path: Path, context: NormalizationContext) -> NormalizationResult:
        presentation = self._presentation(path)
        inventory = self.inventory(path)
        units = []
        warnings: list[str] = []
        if inventory.metadata.get("macros_present"):
            warnings.append("macros were detected and not executed")
        for slide_number, slide in enumerate(presentation.slides, 1):
            parts: list[str] = []
            visual_count = chart_count = 0
            for shape_number, shape in enumerate(self._walk_shapes(slide.shapes), 1):
                label = getattr(shape, "name", f"Object {shape_number}")
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    parts.append(f"### {escape_markdown_cell(label)}\n\n{neutralize_untrusted_text(shape.text)}")
                if getattr(shape, "has_table", False):
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    if rows:
                        parts.append(f"### Table: {escape_markdown_cell(label)}\n\n{markdown_table(rows[0], rows[1:])}")
                if getattr(shape, "has_chart", False):
                    chart_count += 1
                    chart = shape.chart
                    series = [getattr(item, "name", f"Series {index}") for index, item in enumerate(chart.series, 1)]
                    parts.append(
                        f"### Chart: {escape_markdown_cell(label)}\n\n"
                        f"- Type: `{escape_markdown_cell(chart.chart_type)}`\n"
                        f"- Series: {', '.join(escape_markdown_cell(item) for item in series) or 'none'}"
                    )
                if getattr(shape, "shape_type", None) in {6, 13}:
                    visual_count += 1
            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            except Exception:
                notes = ""
            if notes:
                parts.append("### Speaker notes\n\n" + neutralize_untrusted_text(notes))
            if not parts:
                parts.append("> No extractable native slide text was found.")
            visual_state = "REVIEW_REQUIRED" if visual_count or chart_count else "NONE"
            slide_warnings = ["slide rendering is required for layout verification"] if visual_state == "REVIEW_REQUIRED" else []
            units.append(
                unit(
                    locator_kind="slide",
                    locator=f"Slide {slide_number}",
                    filename=f"slide-{slide_number:04d}.md",
                    body=f"## Slide {slide_number}\n\n" + "\n\n".join(parts),
                    visual_review=visual_state,
                    metadata={"slide_number": slide_number, "chart_count": chart_count, "visual_object_count": visual_count, "extraction_mode": "native"},
                    warnings=slide_warnings,
                )
            )
            warnings.extend(f"Slide {slide_number}: {item}" for item in slide_warnings)
        artifacts: list[dict[str, Any]] = []
        if context.profile == "strict-local" and os.getenv("GOV360_OFFICE_RENDERING", "").lower() in {"1", "true", "yes"}:
            artifacts, render_warning = _powerpoint_renders(path, context.output_dir)
            if render_warning:
                warnings.append(render_warning)
        return NormalizationResult(
            self.adapter_id,
            self.adapter_version,
            inventory.source_format,
            "presentation",
            inventory,
            units,
            data_artifacts=artifacts,
            warnings=warnings,
            metrics={"slide_count": len(units), "unit_count": len(units)},
        )


def _chart_metadata(chart: Any) -> dict[str, Any]:
    anchor = getattr(chart, "anchor", None)
    anchor_text = None
    if anchor is not None:
        anchor_text = f"{getattr(anchor, '_from', None)}:{getattr(anchor, 'to', None)}"
    def source_formula(source: Any) -> str | None:
        if source is None:
            return None
        for reference_name in ("numRef", "strRef", "multiLvlStrRef"):
            reference = getattr(source, reference_name, None)
            formula = getattr(reference, "f", None) if reference is not None else None
            if formula:
                return str(formula)
        return None

    def series_title(item: Any) -> str:
        title = getattr(item, "tx", None)
        if title is None:
            return ""
        value = getattr(title, "v", None)
        if value:
            return str(value)
        return str(getattr(getattr(title, "strRef", None), "f", None) or "")

    series: list[dict[str, Any]] = []
    for item in getattr(chart, "ser", ()) or ():
        formulas = [value for value in (
            source_formula(getattr(item, "cat", None)),
            source_formula(getattr(item, "val", None)),
            source_formula(getattr(item, "xVal", None)),
            source_formula(getattr(item, "yVal", None)),
        ) if value]
        series.append({
            "title": series_title(item),
            "categories_formula": source_formula(getattr(item, "cat", None)),
            "values_formula": source_formula(getattr(item, "val", None)) or source_formula(getattr(item, "yVal", None)),
            "source_formulas": formulas,
        })
    try:
        title = str(chart.title or "")
    except Exception:
        title = ""
    return {"type": type(chart).__name__, "title": title, "anchor": anchor_text, "series": series}


class SpreadsheetAdapter(ValidatingAdapter):
    adapter_id = "spreadsheet"
    heavy = True
    capabilities = {"air_gap": True, "native_text": True, "rendering": "local-office", "ocr": False, "tables": True, "charts": True}

    def probe(self, path: Path) -> ProbeResult | None:
        fmt = "xlsm" if path.suffix.lower() == ".xlsm" else "xlsx"
        result = _ooxml_probe(
            path,
            root="xl",
            extensions={".xlsx", ".xlsm"},
            fmt=fmt,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        if result is None:
            return None
        return ProbeResult(self.adapter_id, result.source_format, result.mime_type, result.confidence, result.warnings)

    @staticmethod
    def _workbooks(path: Path):
        try:
            import openpyxl
        except ImportError as exc:
            raise DependencyUnavailableError("openpyxl is required for spreadsheet normalization") from exc
        try:
            formulas = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=True, keep_vba=path.suffix.lower() == ".xlsm")
            cached = openpyxl.load_workbook(path, data_only=True, read_only=False, keep_links=True, keep_vba=path.suffix.lower() == ".xlsm")
            return formulas, cached
        except Exception as exc:
            raise QuarantinedSourceError("spreadsheet package cannot be read") from exc

    def inventory(self, path: Path, context: NormalizationContext | None = None) -> SourceInventory:
        formulas, cached = self._workbooks(path)
        try:
            names = _zip_names(path)
            sheets = []
            for sheet in formulas.worksheets:
                sheets.append({
                    "title": sheet.title,
                    "visibility": sheet.sheet_state,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "table_count": len(sheet.tables),
                    "chart_count": len(sheet._charts),
                    "image_count": len(sheet._images),
                    "merged_range_count": len(sheet.merged_cells.ranges),
                    "comment_count": sum(1 for cell in getattr(sheet, "_cells", {}).values() if cell.comment is not None),
                })
            metadata = {
                "sheet_count": len(sheets),
                "sheets": sheets,
                "chart_count": sum(int(item["chart_count"]) for item in sheets),
                "image_count": sum(int(item["image_count"]) for item in sheets),
                "defined_name_count": len(formulas.defined_names),
                "external_link_count": len(getattr(formulas, "_external_links", ())),
                "macros_present": "xl/vbaProject.bin" in names,
                "calculation_mode": getattr(formulas.calculation, "calcMode", None),
            }
            return SourceInventory("xlsm" if metadata["macros_present"] else "xlsx", "mixed", len(sheets), metadata)
        finally:
            formulas.close()
            cached.close()

    def normalize(self, path: Path, context: NormalizationContext) -> NormalizationResult:
        formulas, cached = self._workbooks(path)
        inventory = self.inventory(path)
        units = []
        warnings: list[str] = []
        if inventory.metadata.get("macros_present"):
            warnings.append("macros were detected and not executed")
        if inventory.metadata.get("external_link_count"):
            warnings.append("external links were not updated; dependent values are UNCERTAIN")
        try:
            for sheet_index, sheet in enumerate(formulas.worksheets, 1):
                cached_sheet = cached[sheet.title]
                used_rows = sorted({cell.row for cell in getattr(sheet, "_cells", {}).values() if cell.value is not None})
                row_groups = [used_rows[offset : offset + context.max_records_per_unit] for offset in range(0, len(used_rows), context.max_records_per_unit)] or [[]]
                sheet_charts = [_chart_metadata(chart) for chart in sheet._charts]
                for part, row_numbers in enumerate(row_groups, 1):
                    first = row_numbers[0] if row_numbers else 1
                    last = row_numbers[-1] if row_numbers else 1
                    cells: list[list[str]] = []
                    formulas_found = 0
                    errors_found = 0
                    for row_number in row_numbers:
                        for cell in sheet[row_number]:
                            if cell.value is None:
                                continue
                            cached_value = cached_sheet[cell.coordinate].value
                            formula = cell.value if cell.data_type == "f" else None
                            formulas_found += int(formula is not None)
                            errors_found += int(cell.data_type == "e")
                            cells.append([
                                cell.coordinate,
                                formula if formula is not None else cell.value,
                                cached_value if formula is not None else cell.value,
                                "yes" if cell.comment is not None else "no",
                            ])
                    table = markdown_table(["Cell", "Formula or native value", "Cached/displayed value", "Comment"], cells)
                    locator = f"Sheet {sheet.title}, cells A{first}:{sheet.cell(last, max(1, sheet.max_column)).coordinate}"
                    review = "REVIEW_REQUIRED" if sheet_charts or sheet._images else "NONE"
                    unit_warnings: list[str] = []
                    if formulas_found:
                        unit_warnings.append("formulas were not recalculated; cached results may be absent or stale")
                    if errors_found:
                        unit_warnings.append("one or more spreadsheet error cells require review")
                    units.append(
                        unit(
                            locator_kind="sheet_range",
                            locator=locator,
                            filename=f"sheet-{sheet_index:03d}-{safe_filename(sheet.title.lower())}-{part:03d}.md",
                            body=(
                                f"## Sheet `{escape_markdown_cell(sheet.title)}` rows {first}-{last}\n\n"
                                f"- Visibility: `{sheet.sheet_state}`\n"
                                f"- Charts: {len(sheet_charts)}\n"
                                f"- Images: {len(sheet._images)}\n"
                                f"- Merged ranges: {len(sheet.merged_cells.ranges)}\n\n"
                                f"{table or '> No non-empty native cells were found.'}\n\n"
                                f"### Chart inventory\n\n{neutralize_untrusted_text(sheet_charts)}"
                            ),
                            visual_review=review,
                            metadata={
                                "sheet_index": sheet_index,
                                "sheet_name": sheet.title,
                                "sheet_visibility": sheet.sheet_state,
                                "row_start": first,
                                "row_end": last,
                                "formula_count": formulas_found,
                                "chart_metadata": sheet_charts,
                                "extraction_mode": "native",
                            },
                            warnings=unit_warnings,
                        )
                    )
                    warnings.extend(f"{locator}: {item}" for item in unit_warnings)
        finally:
            formulas.close()
            cached.close()
        artifacts: list[dict[str, Any]] = []
        if context.profile == "strict-local" and int(inventory.metadata.get("chart_count", 0)) and os.getenv("GOV360_OFFICE_RENDERING", "").lower() in {"1", "true", "yes"}:
            artifacts, render_warning = _excel_chart_renders(path, context.output_dir)
            if render_warning:
                warnings.append(render_warning)
        return NormalizationResult(
            self.adapter_id,
            self.adapter_version,
            inventory.source_format,
            "mixed",
            inventory,
            units,
            data_artifacts=artifacts,
            warnings=warnings,
            metrics={"sheet_count": inventory.unit_count, "unit_count": len(units)},
        )


class OpenDocumentAdapter(ValidatingAdapter):
    adapter_id = "open-document"
    MIMES = {
        "application/vnd.oasis.opendocument.text": ("odt", "document"),
        "application/vnd.oasis.opendocument.presentation": ("odp", "presentation"),
        "application/vnd.oasis.opendocument.spreadsheet": ("ods", "mixed"),
    }

    @staticmethod
    def _kind(path: Path) -> tuple[str, str] | None:
        with path.open("rb") as stream:
            if stream.read(4) != OOXML_SIGNATURE:
                return None
        validate_safe_zip(path)
        with zipfile.ZipFile(path) as package:
            try:
                mime = package.read("mimetype").decode("ascii").strip()
            except (KeyError, UnicodeDecodeError):
                return None
        return OpenDocumentAdapter.MIMES.get(mime)

    def probe(self, path: Path) -> ProbeResult | None:
        kind = self._kind(path)
        if kind is None:
            return None
        fmt, _ = kind
        warning = () if path.suffix.lower() == f".{fmt}" else ("extension does not match OpenDocument signature",)
        mime = {
            "odt": "application/vnd.oasis.opendocument.text",
            "odp": "application/vnd.oasis.opendocument.presentation",
            "ods": "application/vnd.oasis.opendocument.spreadsheet",
        }[fmt]
        return ProbeResult(self.adapter_id, fmt, mime, 1.0, warning)

    @staticmethod
    def _root(path: Path) -> ET.Element:
        try:
            return ET.fromstring(_package_xml(path, "content.xml"))
        except ET.ParseError as exc:
            raise QuarantinedSourceError("OpenDocument content XML is invalid") from exc

    def inventory(self, path: Path, context: NormalizationContext | None = None) -> SourceInventory:
        kind = self._kind(path)
        if kind is None:
            raise QuarantinedSourceError("unsupported OpenDocument package")
        fmt, source_kind = kind
        root = self._root(path)
        target = "page" if fmt == "odp" else "table" if fmt == "ods" else "p"
        count = sum(1 for node in root.iter() if _local_name(node.tag) == target)
        return SourceInventory(fmt, source_kind, max(1, count), {"native_unit_count": count})

    def normalize(self, path: Path, context: NormalizationContext) -> NormalizationResult:
        inventory = self.inventory(path)
        root = self._root(path)
        if inventory.source_format == "odp":
            nodes = [node for node in root.iter() if _local_name(node.tag) == "page"]
            locator_kind, label = "slide", "Slide"
        elif inventory.source_format == "ods":
            nodes = [node for node in root.iter() if _local_name(node.tag) == "table"]
            locator_kind, label = "sheet_range", "Sheet"
        else:
            paragraphs = [node for node in root.iter() if _local_name(node.tag) in {"h", "p"}]
            nodes = [paragraphs[index : index + 100] for index in range(0, len(paragraphs), 100)]
            locator_kind, label = "section", "Section"
        nodes = nodes or [[root]]
        units = []
        for index, node_or_nodes in enumerate(nodes, 1):
            group = node_or_nodes if isinstance(node_or_nodes, list) else [node_or_nodes]
            text = "\n".join("".join(node.itertext()).strip() for node in group if "".join(node.itertext()).strip())
            locator = f"{label} {index}"
            visual = "REVIEW_REQUIRED" if inventory.source_format == "odp" else "NONE"
            units.append(unit(locator_kind=locator_kind, locator=locator, filename=f"{label.lower()}-{index:04d}.md", body=f"## {locator}\n\n{neutralize_untrusted_text(text)}", visual_review=visual, metadata={"extraction_mode": "native"}))
        return NormalizationResult(self.adapter_id, self.adapter_version, inventory.source_format, inventory.source_kind, inventory, units, metrics={"unit_count": len(units)})


class LegacyOfficeAdapter(ValidatingAdapter):
    adapter_id = "legacy-office"
    heavy = True
    FORMATS = {".doc": ("doc", "docx"), ".ppt": ("ppt", "pptx"), ".xls": ("xls", "xlsx")}

    def probe(self, path: Path) -> ProbeResult | None:
        with path.open("rb") as stream:
            if stream.read(8) != OLE_SIGNATURE:
                return None
        fmt = self.FORMATS.get(path.suffix.lower())
        if fmt is None:
            raise QuarantinedSourceError("ambiguous OLE document requires a recognized legacy Office extension")
        return ProbeResult(self.adapter_id, fmt[0], "application/x-ole-storage", 0.99)

    @staticmethod
    def _delegate(path: Path) -> ValidatingAdapter:
        return {".docx": WordProcessingAdapter(), ".pptx": PresentationAdapter(), ".xlsx": SpreadsheetAdapter()}[path.suffix.lower()]

    def _convert(self, path: Path, output: Path) -> Path:
        target_ext = self.FORMATS[path.suffix.lower()][1]
        bridge = output / "legacy-bridge"
        bridge.mkdir(parents=True, exist_ok=True)
        copied = bridge / path.name
        shutil.copy2(path, copied)
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice:
            process = subprocess.run(
                [soffice, "--headless", "--convert-to", target_ext, "--outdir", str(bridge), str(copied)],
                stdin=subprocess.DEVNULL,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=180,
                check=False,
            )
            converted = bridge / f"{copied.stem}.{target_ext}"
            if process.returncode == 0 and converted.exists():
                return converted
        raise DependencyUnavailableError("legacy Office normalization requires a controlled local LibreOffice bridge")

    def inventory(self, path: Path, context: NormalizationContext | None = None) -> SourceInventory:
        fmt = self.FORMATS[path.suffix.lower()][0]
        return SourceInventory(fmt, "document", 1, {"requires_local_conversion": True}, ("legacy Office source requires controlled local conversion",))

    def normalize(self, path: Path, context: NormalizationContext) -> NormalizationResult:
        converted = self._convert(path, context.output_dir)
        delegate = self._delegate(converted)
        result = delegate.normalize(converted, context)
        result.adapter_id = self.adapter_id
        result.adapter_version = self.adapter_version
        result.source_format = self.FORMATS[path.suffix.lower()][0]
        result.inventory = SourceInventory(result.source_format, result.source_kind, result.inventory.unit_count, {**result.inventory.metadata, "converted_locally_to": converted.suffix.lower().lstrip(".")}, tuple(result.inventory.warnings))
        result.warnings.append("legacy Office file was converted locally from an immutable temporary copy")
        shutil.rmtree(converted.parent, ignore_errors=True)
        return result
